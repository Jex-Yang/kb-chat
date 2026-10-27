"""KBChat 文档解析器：pdf/docx/xlsx/csv/txt/md → 纯文本。

白名单 {txt,md,csv,pdf,docx,xlsx}，白名单外抛 ValueError，message 可直接作为 {detail}。
"""
from __future__ import annotations

import csv as csv_mod
import io

ALLOWED_EXTS = {"txt", "md", "csv", "pdf", "docx", "xlsx"}

_ERR_UNSUPPORTED = "不支持的文件类型：.{ext}，仅支持 {allowed}；.doc 请另存为 docx 后再上传"


def check_ext(ext: str) -> None:
    """白名单校验，不通过抛 ValueError（契约 §7）。"""
    ext = ext.lower().lstrip(".")
    if ext not in ALLOWED_EXTS:
        if ext == "doc":
            raise ValueError("不支持 .doc 老格式，请另存为 docx 后再上传")
        raise ValueError(_ERR_UNSUPPORTED.format(ext=ext, allowed="、".join(sorted(ALLOWED_EXTS))))


def parse(ext: str, data: bytes) -> str:
    """按扩展名解析为纯文本。失败抛异常，message 可直接作为 {detail} 返回。"""
    ext = ext.lower().lstrip(".")
    check_ext(ext)
    try:
        if ext in ("txt", "md"):
            return _decode(data)
        if ext == "csv":
            return _parse_csv(data)
        if ext == "pdf":
            return _parse_pdf(data)
        if ext == "docx":
            return _parse_docx(data)
        if ext == "xlsx":
            return _parse_xlsx(data)
    except ValueError:
        raise
    except Exception as exc:  # 解析失败 → detail
        raise ValueError(f"文件解析失败：{type(exc).__name__}: {exc}") from exc
    raise ValueError(_ERR_UNSUPPORTED.format(ext=ext, allowed="、".join(sorted(ALLOWED_EXTS))))


def _decode(data: bytes) -> str:
    for enc in ("utf-8", "gbk"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _parse_csv(data: bytes) -> str:
    """CSV → Markdown 表格文本。"""
    text = _decode(data)
    reader = csv_mod.reader(io.StringIO(text))
    rows = [row for row in reader if any(cell.strip() for cell in row)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    out: list[str] = []
    for i, row in enumerate(rows):
        cells = [c.strip() for c in row] + [""] * (width - len(row))
        out.append("| " + " | ".join(cells) + " |")
        if i == 0:
            out.append("|" + "---|" * width)
    return "\n".join(out)


def _parse_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(text)
    return "\n\n".join(pages)


def _parse_docx(data: bytes) -> str:
    from docx import Document  # python-docx

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    # 按文档体顺序输出段落与表格（简化：先全部段落，再全部表格，保证不丢内容）
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append("| " + " | ".join(cells) + " |")
    return "\n".join(parts)


def _parse_xlsx(data: bytes) -> str:
    """xlsx → 每行转 `| a | b |` 形式，行间换行。"""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"## 工作表：{sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            if row is None or all(v is None or str(v).strip() == "" for v in row):
                continue
            cells = ["" if v is None else str(v).strip() for v in row]
            parts.append("| " + " | ".join(cells) + " |")
    wb.close()
    return "\n".join(parts)
